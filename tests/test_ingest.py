from __future__ import annotations

import importlib.util
import tempfile
import unittest
from pathlib import Path

from secretary.ingest import ingest, load_document


class IngestionTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.store = self.root / "store"

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_txt_preserves_raw_text_and_duplicate_ingest_is_idempotent(self) -> None:
        source = self.root / "note.txt"
        source.write_bytes(b"Prima riga\r\n\r\nSeconda riga\r")

        first = ingest(source, self.store)
        second = ingest(source, self.store)
        document = load_document(first["document"]["document_id"], self.store)

        self.assertFalse(first["duplicate"])
        self.assertTrue(second["duplicate"])
        self.assertEqual(first["document"]["document_id"], second["document"]["document_id"])
        self.assertEqual(document["blocks"][0]["raw_text"], "Prima riga\r\n\r\nSeconda riga\r")
        self.assertEqual(document["blocks"][0]["normalized_text"], "Prima riga\n\nSeconda riga\n")
        self.assertEqual(len(list((self.store / "canonical").glob("doc_*.json"))), 1)

    @unittest.skipUnless(importlib.util.find_spec("fitz"), "PyMuPDF non installato nell'ambiente di test")
    def test_pdf_blocks_include_page_bbox_and_char_range(self) -> None:
        import fitz

        source = self.root / "sample.pdf"
        pdf = fitz.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "Consegna: 12 maggio 2026")
        page.insert_text((72, 120), "Secondo blocco")
        pdf.save(source)
        pdf.close()

        document = ingest(source, self.store)["document"]
        self.assertGreaterEqual(len(document["blocks"]), 2)
        first = document["blocks"][0]
        self.assertEqual(first["source_locator"]["page"], 1)
        self.assertEqual(len(first["source_locator"]["bbox"]), 4)
        self.assertEqual(len(first["source_locator"]["char_range"]), 2)
        self.assertEqual(document["extraction_warnings"], [])

    @unittest.skipUnless(importlib.util.find_spec("docx"), "python-docx non installato nell'ambiente di test")
    def test_docx_preserves_body_order_for_paragraph_and_table(self) -> None:
        from docx import Document

        source = self.root / "sample.docx"
        word = Document()
        word.add_paragraph("Prima del prospetto")
        table = word.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Voce"
        table.cell(0, 1).text = "Valore"
        word.add_paragraph("Dopo il prospetto")
        word.save(source)

        document = ingest(source, self.store)["document"]
        self.assertEqual([item["type"] for item in document["blocks"]], ["paragraph", "table", "paragraph"])
        self.assertEqual([item["source_locator"]["body_index"] for item in document["blocks"]], [1, 2, 3])
        self.assertEqual(document["blocks"][1]["rows"], [["Voce", "Valore"]])
        self.assertEqual(document["extraction_warnings"][0]["code"], "docx_non_body_content_not_extracted")


if __name__ == "__main__":
    unittest.main()
